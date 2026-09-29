from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "Distributed_Physical_Mobile_Device_Lab_Proposal.docx"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Arial"
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_run_font(run, 14 if level == 1 else 12, True, (31, 78, 121))
    return paragraph


def add_paragraph(doc, text=""):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    for run in paragraph.runs:
        set_run_font(run, 11)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run, 11)
    paragraph.paragraph_format.space_after = Pt(3)
    return paragraph


def add_numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    run = paragraph.add_run(text)
    set_run_font(run, 11)
    paragraph.paragraph_format.space_after = Pt(3)
    return paragraph


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)
    return table


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        paragraph = hdr_cells[idx].paragraphs[0]
        run = paragraph.add_run(header)
        set_run_font(run, 10, True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            paragraph = cells[idx].paragraphs[0]
            run = paragraph.add_run(str(value))
            set_run_font(run, 9.5)
    doc.add_paragraph("")
    return table


def build_document():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Design and Implementation of a Distributed Physical Mobile Device Lab for Remote Application Testing")
    set_run_font(run, 17, True, (31, 78, 121))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Project Proposal")
    set_run_font(run, 12, True)

    date_line = doc.add_paragraph()
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_line.add_run("15-Week Project Timeline: August 24 - December 6, 2026")
    set_run_font(run, 10)

    doc.add_paragraph("")

    add_heading(doc, "1. Introduction")
    add_paragraph(
        doc,
        "Mobile application testing requires access to real physical devices because emulators cannot fully reproduce "
        "hardware behavior, vendor Android modifications, screen differences, battery policies, network behavior, and "
        "real mobile browser behavior. A distributed physical mobile device lab allows testers to remotely access Android "
        "phones and run controlled tests without being physically present with the devices."
    )
    add_paragraph(
        doc,
        "This project proposes the design and implementation of a distributed Android device lab using Android Debug Bridge "
        "(ADB), Chrome DevTools Protocol (CDP), a phone-side relay application, a VPS relay server, device screen streaming, "
        "and a desktop controller. The system will support remote device setup, application installation, app launching, "
        "live visual monitoring, native app testing, and real mobile browser automation."
    )

    add_heading(doc, "2. Problem Statement")
    add_paragraph(
        doc,
        "Developers and QA teams often have limited access to physical Android devices for testing. Existing cloud device "
        "labs provide useful testing infrastructure, but they can be expensive, less customizable, and limited for persistent "
        "private device sessions. Smaller teams, research groups, and organizations may need a self-hosted way to organize "
        "their own Android devices into a remotely accessible testing lab."
    )
    add_paragraph(
        doc,
        "The problem is to design a practical system that allows privately owned Android devices to be remotely connected, "
        "monitored, streamed, and controlled for repeatable mobile application and browser testing through authorized ADB "
        "and CDP access."
    )

    add_heading(doc, "3. Project Aim and Objectives")
    add_paragraph(doc, "Aim:")
    add_paragraph(
        doc,
        "To design and implement a distributed physical mobile device lab for remote Android application and browser testing."
    )
    add_paragraph(doc, "Objectives:")
    objectives = [
        "Design a distributed architecture for remotely controlled Android devices.",
        "Implement a phone-side relay app for status reporting and device connectivity.",
        "Implement a VPS relay server for connecting remote devices to a controller.",
        "Implement a desktop controller for ADB setup and device bootstrap.",
        "Enable remote ADB-based control for app installation, launch, screenshots, logs, and testing.",
        "Enable device screen streaming for live remote viewing and interactive test observation.",
        "Enable Chrome/CDP-based real mobile browser automation.",
        "Evaluate reliability across local, remote, foreground, background, screen-off, and reconnect scenarios.",
    ]
    for item in objectives:
        add_numbered(doc, item)

    add_heading(doc, "4. Research Questions")
    questions = [
        "How can physical Android devices be connected into a distributed remote testing lab?",
        "How can ADB support remote application installation, launching, control, monitoring, and artifact collection?",
        "How can a relay architecture reduce dependency on direct physical access after initial setup?",
        "How can screen streaming improve remote visibility and interactive testing?",
        "How can CDP extend the system to support real mobile browser automation?",
        "What reliability and security limitations affect remote physical device testing?",
    ]
    for item in questions:
        add_numbered(doc, item)

    add_heading(doc, "5. Related Work and Research Gap")
    add_paragraph(
        doc,
        "Existing services such as Firebase Test Lab and AWS Device Farm provide access to real mobile devices for application "
        "testing. Android Device Streaming also provides remote access to selected physical Android devices from Android Studio. "
        "These services demonstrate the value of physical devices in mobile QA."
    )
    add_paragraph(
        doc,
        "However, centralized cloud labs may not fully address use cases where an organization wants to operate its own private "
        "distributed device pool, keep persistent device sessions, customize device setup, or combine ADB-level device control "
        "with CDP-level browser automation. The research gap is a self-hosted distributed Android device lab that combines "
        "private physical devices, ADB-based app testing, live screen streaming, remote relay connectivity, desktop bootstrap "
        "tooling, and real Chrome/CDP browser automation."
    )

    add_heading(doc, "6. Proposed System Architecture")
    add_paragraph(doc, "The proposed system contains five main components:")
    components = [
        "Android Relay App: runs on the physical Android device, maintains a foreground service, reports device status, connects outward to the relay server, and probes local ADB availability.",
        "VPS Relay Server: receives phone relay connections, bridges controller traffic to remote devices, and provides a controlled path for remote ADB access.",
        "Desktop Controller: detects USB-connected Android devices, installs or updates the relay app, starts the relay app, enables ADB TCP mode, and reports setup status.",
        "Screen Streaming Layer: provides live device viewing for remote observation and interactive testing using ADB-based capture or a streaming tool such as scrcpy.",
        "Automation Controller: uses ADB for native app operations and CDP for Chrome browser automation, then collects screenshots, logs, recordings, and test results.",
    ]
    for item in components:
        add_bullet(doc, item)

    add_paragraph(doc, "Architecture Sketch:")
    add_code_block(
        doc,
        """
                     +-----------------------------+
                     |       Tester / QA User      |
                     | Desktop Controller + Viewer |
                     +--------------+--------------+
                                    |
                         commands + screen stream
                                    |
                     +--------------v--------------+
                     |     Automation Controller   |
                     |   ADB + CDP + Test Runner   |
                     +--------------+--------------+
                                    |
                                    | remote control path
                                    |
                     +--------------v--------------+
                     |        VPS Relay Server     |
                     |   ADB tunnel + stream path  |
                     +--------------+--------------+
                                    ^
                                    |
                          outbound phone connection
                                    |
+-----------------------------------+-----------------------------------+
|                          Physical Android Device                     |
|                                                                       |
|  +-----------------------+       +---------------------------------+  |
|  |    All in Relay App   | <---> | Local ADB TCP / Device State    |  |
|  | Status + connection   |       | 127.0.0.1:5555                  |  |
|  +-----------------------+       +---------------------------------+  |
|                                                                       |
|  +-----------------------+       +---------------------------------+  |
|  | App Under Test        |       | Chrome Mobile Browser           |  |
|  | ADB/UI automation     |       | CDP browser automation          |  |
|  +-----------------------+       +---------------------------------+  |
|                                                                       |
|  +---------------------------------------------------------------+    |
|  | Live Screen Stream / Screenshot / Recording Capture           |    |
|  +---------------------------------------------------------------+    |
+-----------------------------------------------------------------------+
""".strip(),
    )

    add_heading(doc, "7. Methodology and Implementation Approach")
    add_paragraph(
        doc,
        "The project will follow an iterative prototype methodology. The first stage will implement direct USB-based ADB "
        "control and local testing. The second stage will add the Android relay application and VPS relay server. The third "
        "stage will integrate remote ADB tunneling and screen streaming. The final stage will expand testing support, evaluate "
        "reliability, and document the results."
    )
    implementation_items = [
        "ADB device detection and setup checks",
        "relay app installation and startup",
        "ADB TCP bootstrap after trusted setup",
        "phone-to-server relay connection",
        "remote ADB tunneling through the VPS",
        "live screen streaming and screenshot capture",
        "app installation, app launch, and foreground UI control",
        "Chrome/CDP browser connection and browser test execution",
        "test result storage, logs, recordings, and reliability reports",
    ]
    for item in implementation_items:
        add_bullet(doc, item)

    add_heading(doc, "8. System Evaluation")
    add_table(
        doc,
        ["Evaluation Area", "Measurement"],
        [
            ["Device detection", "Whether the desktop controller detects authorized Android devices."],
            ["Relay app setup", "Whether the relay app installs and starts successfully."],
            ["ADB bootstrap", "Whether ADB TCP mode can be enabled after trusted setup."],
            ["Remote relay", "Whether the phone can connect to the VPS relay."],
            ["Remote ADB", "Whether the controller can access the phone through the relay."],
            ["Screen streaming", "Whether the tester can view the live device screen remotely with acceptable latency and stability."],
            ["App testing", "Whether apps can be installed, launched, controlled, and observed."],
            ["Browser testing", "Whether Chrome can be controlled through CDP."],
            ["Data collection", "Whether screenshots, logs, recordings, and test results are captured."],
            ["Reliability", "Behavior during foreground, background, screen-off, and reconnect scenarios."],
            ["Security", "Whether access is limited to authorized setup and controlled relay paths."],
        ],
    )

    add_heading(doc, "9. Project Scope and Limitations")
    add_paragraph(doc, "Scope:")
    scope_items = [
        "Android physical device setup and management",
        "ADB-based remote control",
        "phone relay communication",
        "VPS relay connectivity",
        "desktop bootstrap tooling",
        "native Android app testing",
        "real Chrome browser automation through CDP",
        "live screen streaming, screenshots, logs, recordings, and result reporting",
    ]
    for item in scope_items:
        add_bullet(doc, item)
    add_paragraph(doc, "Limitations:")
    limitation_items = [
        "The project requires initial owner-authorized ADB setup.",
        "The system will not bypass lock screens.",
        "Native app UI testing requires the target app to be in the foreground.",
        "Android foreground services may show system indicators.",
        "Screen streaming quality may depend on bandwidth, device performance, and relay latency.",
        "Some secure screens may block screenshots or recordings.",
        "The first implementation may use basic token-based relay security before stronger TLS/mTLS hardening.",
        "iOS support is outside the main project scope.",
    ]
    for item in limitation_items:
        add_bullet(doc, item)

    add_heading(doc, "10. Expected Results and Contributions")
    expected_items = [
        "A working Android relay application.",
        "A desktop controller for ADB bootstrap.",
        "A VPS relay server for remote device connectivity.",
        "A prototype remote ADB tunnel.",
        "A live device screen streaming feature.",
        "A native Android app testing workflow.",
        "A Chrome/CDP browser automation workflow.",
        "Captured test artifacts such as logs, screenshots, recordings, and result files.",
        "A documented architecture for a private distributed Android device lab.",
    ]
    for item in expected_items:
        add_numbered(doc, item)
    add_paragraph(
        doc,
        "The main contribution is a practical prototype showing how privately controlled physical Android devices can be "
        "organized into a remote testing lab using ADB, CDP, screen streaming, and relay-based connectivity."
    )

    add_heading(doc, "11. Project Timeline")
    add_paragraph(doc, "15 Weeks: August 24 - December 6, 2026")
    add_table(
        doc,
        ["Week", "Dates", "Activities", "Expected Output"],
        [
            ["1", "Aug 24 - Aug 30", "Define scope, requirements, and system goals.", "Approved proposal direction."],
            ["2", "Aug 31 - Sep 6", "Set up Node.js, Android, ADB, and project structure.", "Working development environment."],
            ["3", "Sep 7 - Sep 13", "Implement ADB preflight checks and device detection.", "ADB device detection scripts."],
            ["4", "Sep 14 - Sep 20", "Build Chrome/CDP baseline automation.", "Browser automation proof of concept."],
            ["5", "Sep 21 - Sep 27", "Test foreground, background, and screen-off browser behavior.", "Browser behavior test results."],
            ["6", "Sep 28 - Oct 4", "Build Android relay app prototype.", "Phone-side relay app."],
            ["7", "Oct 5 - Oct 11", "Add local status API, foreground service, and ADB probe.", "Phone status/control API."],
            ["8", "Oct 12 - Oct 18", "Build VPS relay server and heartbeat system.", "Phone-to-server relay connection."],
            ["9", "Oct 19 - Oct 25", "Implement ADB byte tunnel through VPS relay.", "Remote ADB tunnel prototype."],
            ["10", "Oct 26 - Nov 1", "Build desktop controller for device setup.", "Desktop ADB setup app."],
            ["11", "Nov 2 - Nov 8", "Add app install, launch, screenshots, logs, and initial screen streaming.", "Native app workflow with visual monitoring."],
            ["12", "Nov 9 - Nov 15", "Improve screen streaming and integrate Chrome/CDP through remote access.", "Remote browser testing and stable device viewing."],
            ["13", "Nov 16 - Nov 22", "Improve reliability, reconnection, error handling, stream stability, and result storage.", "Stabilized prototype."],
            ["14", "Nov 23 - Nov 29", "Evaluate performance, limitations, usability, and security.", "Evaluation results and analysis."],
            ["15", "Nov 30 - Dec 6", "Prepare final report, diagrams, demo script, and presentation.", "Final project submission."],
        ],
    )

    add_heading(doc, "12. Conclusion")
    add_paragraph(
        doc,
        "This project proposes a distributed physical Android device lab for remote application testing. By combining ADB, CDP, "
        "a phone relay app, a VPS relay server, device screen streaming, and a desktop controller, the system will allow real "
        "Android devices to be prepared, connected, viewed, controlled, and tested remotely."
    )
    add_paragraph(
        doc,
        "The expected outcome is a working prototype that demonstrates remote native app testing and real mobile browser "
        "automation on physical Android devices."
    )

    add_heading(doc, "References")
    references = [
        "Android Developers. Android Debug Bridge. https://developer.android.com/tools/adb",
        "Chrome DevTools. Chrome DevTools Protocol. https://chromedevtools.github.io/",
        "Google Firebase. Firebase Test Lab. https://firebase.google.com/docs/test-lab",
        "Amazon Web Services. AWS Device Farm. https://aws.amazon.com/device-farm/",
        "Android Developers. Android Device Streaming. https://developer.android.com/studio/run/android-device-streaming",
    ]
    for item in references:
        add_numbered(doc, item)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
